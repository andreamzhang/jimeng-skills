#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
分镜表Excel生成脚本模板
支持两种模式：
- 模式A：从即梦分镜提示词(02-jimeng-prompts.md)解析已填写的7字段内容
- 模式B：从导演讲戏本(01-director-analysis.md)按剧情点生成填写模板（7字段留空待填）
"""

import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ========== 样式定义 ==========
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

ZEBRA_FILL_1 = PatternFill(fill_type="solid", fgColor="FFFFFF")
ZEBRA_FILL_2 = PatternFill(fill_type="solid", fgColor="F7F9FC")

DATA_FONT = Font(size=10, color="000000")
DATA_ALIGN_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
DATA_ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="D9DEE7"),
    right=Side(style="thin", color="D9DEE7"),
    top=Side(style="thin", color="D9DEE7"),
    bottom=Side(style="thin", color="D9DEE7"),
)


# ========== 模式A：解析即梦分镜提示词 ==========
def parse_jimeng_prompts(text):
    """解析即梦分镜提示词，提取母镜头和子镜头的7字段内容"""
    shots = []
    
    # 按"## 母镜头"分割
    mother_shots = re.split(r'(?=##\s+母镜头)', text)
    
    for ms in mother_shots:
        if not ms.strip() or '母镜头' not in ms:
            continue
        
        # 解析母镜头头部
        ms_header = re.search(r'##\s+(母镜头\S+)\s*（总时长：(\d+)秒）\s*——\s*(.+?)\n', ms)
        if not ms_header:
            continue
        ms_id = ms_header.group(1)
        ms_duration = ms_header.group(2)
        ms_title = ms_header.group(3).strip()
        
        # 场景类型
        scene_type_match = re.search(r'\*\*场景类型\*\*：(.+?)\n', ms)
        scene_type = scene_type_match.group(1).strip() if scene_type_match else ""
        
        # 分镜过渡
        transition_match = re.search(r'\*\*分镜过渡\*\*：(.+?)\n\n', ms, re.DOTALL)
        transition = transition_match.group(1).strip() if transition_match else ""
        
        # 解析子镜头
        sub_shots = re.split(r'(?=###\s+子镜头\d+)', ms)
        
        for ss in sub_shots:
            if not ss.strip() or '子镜头' not in ss:
                continue
            
            ss_header = re.search(r'###\s+子镜头(\d+)（时长：([\d.]+)秒）：', ss)
            if not ss_header:
                continue
            ss_id = ss_header.group(1)
            ss_duration = ss_header.group(2)
            
            # 提取7个字段
            fields = {}
            for field_name in ['景别', '运镜', '机位', '画面', '微表情', '台词', '音效']:
                pattern = rf'【{field_name}】(.+?)(?=\n【|$)'
                match = re.search(pattern, ss, re.DOTALL)
                fields[field_name] = match.group(1).strip() if match else ""
            
            shots.append({
                '母镜头编号': ms_id,
                '母镜头标题': ms_title,
                '母镜头总时长': ms_duration,
                '场景类型': scene_type,
                '子镜头编号': ss_id,
                '子镜头时长': ss_duration,
                '景别': fields.get('景别', ''),
                '运镜': fields.get('运镜', ''),
                '机位': fields.get('机位', ''),
                '画面': fields.get('画面', ''),
                '微表情': fields.get('微表情', ''),
                '台词': fields.get('台词', ''),
                '音效': fields.get('音效', '')
            })
    
    return shots


def generate_mode_a(input_path, output_path):
    """模式A：从即梦分镜提示词生成完整分镜表"""
    with open(input_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    shots = parse_jimeng_prompts(content)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "分镜表"
    
    # 表头
    headers = [
        "母镜头编号", "母镜头标题", "母镜头总时长(秒)", "场景类型", "子镜头编号",
        "子镜头时长(秒)", "景别", "运镜", "机位", "画面", "微表情", "台词", "音效"
    ]
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    
    # 列宽
    col_widths = [12, 20, 14, 12, 12, 14, 15, 25, 25, 50, 35, 50, 35]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    
    # 填充数据
    for row_idx, shot in enumerate(shots, 2):
        values = [
            shot['母镜头编号'], shot['母镜头标题'], shot['母镜头总时长'],
            shot['场景类型'], shot['子镜头编号'], shot['子镜头时长'],
            shot['景别'], shot['运镜'], shot['机位'],
            shot['画面'], shot['微表情'], shot['台词'], shot['音效']
        ]
        
        fill = ZEBRA_FILL_1 if (row_idx - 2) % 2 == 0 else ZEBRA_FILL_2
        
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = DATA_FONT
            cell.alignment = DATA_ALIGN_LEFT
            cell.border = THIN_BORDER
            cell.fill = fill
    
    # 行高
    ws.row_dimensions[1].height = 35
    for r in range(2, len(shots) + 2):
        ws.row_dimensions[r].height = 80
    
    ws.freeze_panes = "A2"
    wb.save(output_path)
    return len(shots)


# ========== 模式B：从导演讲戏本生成填写模板 ==========
def parse_director_analysis(text):
    """解析导演讲戏本，提取剧情点（每个剧情点 = 一个母镜头）"""
    plot_points = []
    
    # 按 "## P" 分割剧情点
    sections = re.split(r'(?=##\s+P\d+)', text)
    
    for sec in sections:
        if not sec.strip() or not re.match(r'##\s+P\d+', sec):
            continue
        
        # 剧情点编号和标题
        header = re.search(r'##\s+(P\d+)\s+(.+)', sec)
        if not header:
            continue
        pp_id = header.group(1)
        pp_title = header.group(2).strip()
        
        # 预估时长
        duration_match = re.search(r'预估时长[：:]\s*(\d+)秒', sec)
        duration = duration_match.group(1) if duration_match else ""
        
        # 情绪基调
        mood_match = re.search(r'情绪基调[：:]\s*(.+)', sec)
        mood = mood_match.group(1).strip() if mood_match else ""
        
        # 导演阐述
        desc_match = re.search(r'\*\*导演阐述\*\*[：:]\s*(.+?)(?=\n\*\*|\Z)', sec, re.DOTALL)
        description = desc_match.group(1).strip() if desc_match else ""
        
        plot_points.append({
            '母镜头编号': pp_id,
            '母镜头标题': pp_title,
            '情绪基调': mood,
            '预估时长': duration,
            '导演阐述': description
        })
    
    return plot_points


def generate_mode_b(director_analysis_path, output_path):
    """模式B：从导演讲戏本按剧情点生成填写模板"""
    with open(director_analysis_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    plot_points = parse_director_analysis(content)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "分镜表"
    
    # 表头
    headers = [
        "母镜头编号", "母镜头标题", "情绪基调", "预估时长(秒)", "导演阐述",
        "景别", "运镜", "机位", "画面", "微表情", "台词", "音效",
        "备注"
    ]
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    
    # 列宽
    col_widths = [12, 20, 12, 12, 50, 12, 30, 30, 40, 30, 40, 30, 20]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    
    # 填充数据
    for row_idx, pp in enumerate(plot_points, 2):
        values = [
            pp['母镜头编号'], pp['母镜头标题'], pp['情绪基调'], pp['预估时长'],
            pp['导演阐述'], "", "", "", "", "", "", "", ""
        ]
        
        fill = ZEBRA_FILL_1 if (row_idx - 2) % 2 == 0 else ZEBRA_FILL_2
        
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = DATA_FONT
            cell.alignment = DATA_ALIGN_LEFT
            cell.border = THIN_BORDER
            cell.fill = fill
    
    # 行高
    ws.row_dimensions[1].height = 35
    for r in range(2, len(plot_points) + 2):
        ws.row_dimensions[r].height = 45
    
    ws.freeze_panes = "A2"
    wb.save(output_path)
    return len(plot_points)


# ========== 主函数 ==========
def main():
    """
    使用方式（由技能调用时传入参数）：
    python generate_table_template.py <mode> <input_path> [output_path]
    
    mode: A 或 B
    input_path: 
        - 模式A: 02-jimeng-prompts.md 路径
        - 模式B: 01-director-analysis.md 路径
    output_path: 输出Excel路径（可选）
    """
    if len(sys.argv) < 3:
        print("参数不足，请参考技能说明使用")
        sys.exit(1)
    
    mode = sys.argv[1]
    input_path = sys.argv[2]
    
    if len(sys.argv) >= 4:
        output_path = sys.argv[3]
    else:
        output_path = "分镜表.xlsx"
    
    if mode == "A":
        count = generate_mode_a(input_path, output_path)
        print(f"模式A完成：共 {count} 个子镜头，保存至 {output_path}")
    elif mode == "B":
        count = generate_mode_b(input_path, output_path)
        print(f"模式B完成：共 {count} 个剧情点（母镜头），保存至 {output_path}")
    else:
        print(f"未知模式：{mode}，请使用 A 或 B")
        sys.exit(1)


if __name__ == "__main__":
    main()
